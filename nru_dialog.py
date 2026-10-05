import os
import tempfile
from datetime import datetime

from qgis.PyQt.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QGridLayout,
    QGroupBox, QLabel, QLineEdit, QPushButton,
    QRadioButton, QFileDialog,
    QProgressBar, QMessageBox, QFrame, QTextEdit
)
from qgis.PyQt.QtCore import QSize
from qgis.PyQt.QtGui import QPainter, QColor, QFont
from qgis.core import (
    QgsVectorLayer, QgsGeometry, QgsPointXY,
    QgsFeature, QgsMapSettings, QgsRectangle,
    QgsMapRendererSequentialJob,
    QgsSingleSymbolRenderer,
    QgsFillSymbol, QgsLineSymbol, QgsMarkerSymbol
)
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_COLOR_INDEX


class NRUDialog(QDialog):

    def __init__(self, iface):
        super().__init__(iface.mainWindow())
        self.iface = iface
        self.setWindowTitle("NRU Generator")
        self.setMinimumWidth(620)
        self.resize(650, 700)
        self._build_ui()

    def _build_ui(self):
        main = QVBoxLayout(self)
        main.setSpacing(8)

        grp_a = QGroupBox("A. Informations du demandeur")
        grid = QGridLayout(grp_a)
        grid.setVerticalSpacing(6)
        self.e_nom     = QLineEdit(); self.e_nom.setPlaceholderText("Nom et Prenom")
        self.e_adresse = QLineEdit(); self.e_adresse.setPlaceholderText("Adresse")
        self.e_ref     = QLineEdit(); self.e_ref.setPlaceholderText("Reference fonciere  ex: TF 12345/C")
        self.e_num     = QLineEdit(); self.e_num.setPlaceholderText("N de la demande  ex: 2025/NRU/042")
        self.e_date    = QLineEdit(); self.e_date.setPlaceholderText("Date de la demande  ex: 21/06/2026")
        for i, (lbl, w) in enumerate([
            ("Nom :", self.e_nom), ("Adresse :", self.e_adresse),
            ("Reference :", self.e_ref), ("N demande :", self.e_num), ("Date :", self.e_date)
        ]):
            grid.addWidget(QLabel(lbl), i, 0)
            grid.addWidget(w, i, 1)
        main.addWidget(grp_a)

        grp_b = QGroupBox("B. Localisation de la parcelle")
        vb = QVBoxLayout(grp_b)
        vb.setSpacing(6)
        self.radio_xy  = QRadioButton("Coordonnees X, Y manuelles")
        self.radio_txt = QRadioButton("Fichier texte (.txt)")
        self.radio_shp = QRadioButton("Shapefile (.shp)")
        self.radio_xy.setChecked(True)
        vb.addWidget(self.radio_xy)
        vb.addWidget(self.radio_txt)
        vb.addWidget(self.radio_shp)

        self.pan_xy = QFrame()
        hxy = QHBoxLayout(self.pan_xy)
        self.e_x = QLineEdit(); self.e_x.setPlaceholderText("X  ex: 301603")
        self.e_y = QLineEdit(); self.e_y.setPlaceholderText("Y  ex: 329176")
        hxy.addWidget(QLabel("X:")); hxy.addWidget(self.e_x)
        hxy.addWidget(QLabel("Y:")); hxy.addWidget(self.e_y)
        vb.addWidget(self.pan_xy)

        self.pan_txt = QFrame()
        htxt = QHBoxLayout(self.pan_txt)
        self.e_txt = QLineEdit(); self.e_txt.setReadOnly(True)
        btn_txt = QPushButton("Parcourir...")
        btn_txt.clicked.connect(lambda: self._browse(self.e_txt, "Texte (*.txt)"))
        htxt.addWidget(self.e_txt); htxt.addWidget(btn_txt)
        self.pan_txt.setVisible(False)
        vb.addWidget(self.pan_txt)

        self.pan_shp = QFrame()
        hshp = QHBoxLayout(self.pan_shp)
        self.e_shp = QLineEdit(); self.e_shp.setReadOnly(True)
        btn_shp = QPushButton("Parcourir...")
        btn_shp.clicked.connect(lambda: self._browse(self.e_shp, "Shapefile (*.shp)"))
        hshp.addWidget(self.e_shp); hshp.addWidget(btn_shp)
        self.pan_shp.setVisible(False)
        vb.addWidget(self.pan_shp)

        self.radio_xy.toggled.connect(lambda c: (
            self.pan_xy.setVisible(c), self.pan_txt.setVisible(False), self.pan_shp.setVisible(False)))
        self.radio_txt.toggled.connect(lambda c: (
            self.pan_txt.setVisible(c), self.pan_xy.setVisible(False), self.pan_shp.setVisible(False)))
        self.radio_shp.toggled.connect(lambda c: (
            self.pan_shp.setVisible(c), self.pan_xy.setVisible(False), self.pan_txt.setVisible(False)))
        main.addWidget(grp_b)

        grp_c = QGroupBox("C. Donnees SIG")
        gc = QGridLayout(grp_c)
        gc.setVerticalSpacing(6)
        self.e_gpkg = QLineEdit(); self.e_gpkg.setReadOnly(True)
        btn_gpkg = QPushButton("Parcourir...")
        btn_gpkg.clicked.connect(lambda: self._browse(self.e_gpkg, "GeoPackage (*.gpkg)"))
        gc.addWidget(QLabel("Casa_urba.gpkg :"), 0, 0)
        gc.addWidget(self.e_gpkg, 0, 1)
        gc.addWidget(btn_gpkg, 0, 2)
        self.e_sortie = QLineEdit(); self.e_sortie.setReadOnly(True)
        btn_sortie = QPushButton("Parcourir...")
        btn_sortie.clicked.connect(self._browse_folder)
        gc.addWidget(QLabel("Dossier sortie :"), 1, 0)
        gc.addWidget(self.e_sortie, 1, 1)
        gc.addWidget(btn_sortie, 1, 2)
        main.addWidget(grp_c)

        self.progress = QProgressBar()
        self.progress.setValue(0)
        main.addWidget(self.progress)

        self.log_box = QTextEdit()
        self.log_box.setReadOnly(True)
        self.log_box.setFixedHeight(70)
        main.addWidget(self.log_box)

        hbtn = QHBoxLayout()
        btn_gen = QPushButton("Generer la NRU")
        btn_gen.setMinimumHeight(42)
        btn_gen.clicked.connect(self.generer)
        btn_close = QPushButton("Fermer")
        btn_close.setMinimumHeight(42)
        btn_close.clicked.connect(self.reject)
        hbtn.addStretch()
        hbtn.addWidget(btn_close)
        hbtn.addWidget(btn_gen)
        main.addLayout(hbtn)

    def _browse(self, edit, filtre):
        path, _ = QFileDialog.getOpenFileName(self, "Choisir un fichier", "", filtre)
        if path:
            edit.setText(path)

    def _browse_folder(self):
        d = QFileDialog.getExistingDirectory(self, "Choisir le dossier de sortie")
        if d:
            self.e_sortie.setText(d)

    def _log(self, msg):
        self.log_box.append(msg)

    def _get_geom(self):
        if self.radio_xy.isChecked():
            x = float(self.e_x.text().strip().replace(",", "."))
            y = float(self.e_y.text().strip().replace(",", "."))
            return QgsGeometry.fromPointXY(QgsPointXY(x, y))
        elif self.radio_txt.isChecked():
            points = []
            with open(self.e_txt.text(), encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if not line or line.upper().startswith("X"):
                        continue
                    sep = ";" if ";" in line else ","
                    parts = line.split(sep)
                    x = float(parts[0].strip().replace(",", "."))
                    y = float(parts[1].strip().replace(",", "."))
                    points.append(QgsPointXY(x, y))
            if len(points) >= 3:
                if points[0] != points[-1]:
                    points.append(points[0])
                return QgsGeometry.fromPolygonXY([points])
            return QgsGeometry.fromMultiPointXY(points)
        else:
            layer = QgsVectorLayer(self.e_shp.text(), "parcelle", "ogr")
            return list(layer.getFeatures())[0].geometry()

    def _get_modele(self, geom, limite_du):
        du_geom = QgsGeometry()
        for f in limite_du.getFeatures():
            g = f.geometry()
            if g and not g.isEmpty():
                if du_geom.isEmpty():
                    du_geom = g
                else:
                    du_geom = du_geom.combine(g)
        if du_geom.isEmpty() or not geom.intersects(du_geom):
            return "NRU-3"
        if du_geom.contains(geom):
            return "NRU-1"
        return "NRU-2"

    def _extraire_infos(self, geom, couches):
        infos = {
            "zones": [], "equipements": [],
            "commune": "", "province": "",
            "du_nom": "", "du_type": ""
        }
        for feat in couches["zones"].getFeatures():
            if feat.geometry() and geom.intersects(feat.geometry()):
                infos["zones"].append({
                    "zone":    feat["ZONE"]    or "",
                    "secteur": feat["SECTEUR"] or ""
                })
        for feat in couches["equipements"].getFeatures():
            if feat.geometry() and geom.intersects(feat.geometry()):
                desc = feat["DESCRIPTIO"] or ""
                if desc and desc not in infos["equipements"]:
                    infos["equipements"].append(desc)
        for feat in couches["communes"].getFeatures():
            if feat.geometry() and geom.intersects(feat.geometry()):
                infos["commune"]  = feat["com_fr"]  or ""
                infos["province"] = feat["prov_fr"] or ""
                break
        for feat in couches["limite_du"].getFeatures():
            if feat.geometry() and geom.intersects(feat.geometry()):
                infos["du_nom"]  = feat["Nom"] or "Plan d'Amenagement"
                infos["du_type"] = feat["PA"]  or "Plan d'Amenagement"
                break
        return infos

    def _export_png(self, geom, couches, path):
        style_communes = QgsFillSymbol.createSimple({
            "color": "0,0,0,0", "outline_color": "100,100,100,255", "outline_width": "0.8"})
        style_zones = QgsFillSymbol.createSimple({
            "color": "200,230,200,180", "outline_color": "80,120,80,255", "outline_width": "0.5"})
        style_equip = QgsFillSymbol.createSimple({
            "color": "255,200,100,180", "outline_color": "200,100,0,255", "outline_width": "0.5"})
        style_limite = QgsLineSymbol.createSimple({
            "color": "200,0,0,255", "width": "1.5"})

        couches["communes"].setRenderer(QgsSingleSymbolRenderer(style_communes))
        couches["zones"].setRenderer(QgsSingleSymbolRenderer(style_zones))
        couches["equipements"].setRenderer(QgsSingleSymbolRenderer(style_equip))
        couches["limite_du"].setRenderer(QgsSingleSymbolRenderer(style_limite))

        # Le type de couche memoire ET le type de symbole doivent correspondre
        # au type de geometrie de la parcelle, sinon QGIS n'affiche rien.
        if geom.type() == 2:
            type_g = "Polygon"
        elif geom.type() == 1:
            type_g = "Linestring"
        else:
            type_g = "Point"

        tmp = QgsVectorLayer(type_g + "?crs=EPSG:26191", "Parcelle", "memory")
        feat = QgsFeature()
        feat.setGeometry(geom)
        tmp.dataProvider().addFeature(feat)
        tmp.updateExtents()

        if type_g == "Polygon":
            style_parcelle = QgsFillSymbol.createSimple({
                "color": "255,0,0,70",
                "outline_color": "255,0,0,255",
                "outline_width": "1.4"})
        elif type_g == "Linestring":
            style_parcelle = QgsLineSymbol.createSimple({
                "color": "255,0,0,255", "width": "1.8"})
        else:
            style_parcelle = QgsMarkerSymbol.createSimple({
                "color": "255,0,0,255", "size": "8",
                "outline_color": "255,255,255,255", "outline_width": "1"})
        tmp.setRenderer(QgsSingleSymbolRenderer(style_parcelle))

        bbox = geom.boundingBox()
        marge = 5000
        extent = QgsRectangle(
            bbox.xMinimum() - marge, bbox.yMinimum() - marge,
            bbox.xMaximum() + marge, bbox.yMaximum() + marge)

        settings = QgsMapSettings()
        # La 1ere couche de la liste est dessinee AU-DESSUS des suivantes.
        # La parcelle doit donc passer en premier pour ne pas etre cachee
        # sous les polygones zones/equipements.
        settings.setLayers([
            tmp, couches["limite_du"],
            couches["equipements"], couches["zones"], couches["communes"]])
        settings.setExtent(extent)
        settings.setOutputSize(QSize(1200, 900))
        settings.setBackgroundColor(QColor(230, 240, 255))
        settings.setFlag(QgsMapSettings.Antialiasing, True)

        job = QgsMapRendererSequentialJob(settings)
        job.start()
        job.waitForFinished()
        img = job.renderedImage()

        p = QPainter(img)

        # --- Bandeau de titre ---
        p.setFont(QFont("Arial", 13, QFont.Bold))
        p.setPen(QColor(30, 30, 80))
        p.fillRect(0, img.height() - 50, img.width(), 50, QColor(255, 255, 255, 220))
        p.drawText(10, img.height() - 15,
                   "Plan de Situation - Document d'Urbanisme de Casablanca")
        p.drawRect(2, 2, img.width() - 4, img.height() - 4)

        # --- Legende ---
        legend_items = [
            (QColor(100, 100, 100), "Limite Commune"),
            (QColor(150, 190, 150), "Zones"),
            (QColor(230, 160, 60),  "Equipements"),
            (QColor(200, 0, 0),     "Limite DU"),
            (QColor(255, 0, 0),     "Parcelle"),
        ]
        box_w, box_h = 190, 24 * len(legend_items) + 30
        legend_x = img.width() - box_w - 15
        legend_y = 15

        p.fillRect(legend_x, legend_y, box_w, box_h, QColor(255, 255, 255, 235))
        p.setPen(QColor(50, 50, 50))
        p.drawRect(legend_x, legend_y, box_w, box_h)

        p.setFont(QFont("Arial", 10, QFont.Bold))
        p.setPen(QColor(30, 30, 80))
        p.drawText(legend_x + 10, legend_y + 20, "Legende")

        p.setFont(QFont("Arial", 9))
        yy = legend_y + 40
        for color, label in legend_items:
            p.setBrush(color)
            p.setPen(QColor(60, 60, 60))
            p.drawRect(legend_x + 10, yy - 11, 16, 16)
            p.setPen(QColor(20, 20, 20))
            p.drawText(legend_x + 34, yy, label)
            yy += 24

        p.end()

        ok = img.save(path, "PNG")
        if not ok:
            raise Exception("Impossible de sauvegarder le plan PNG : " + path)

    def _remplacer_para(self, para, subs):
        full = "".join(r.text for r in para.runs)
        new = full
        for k, v in subs.items():
            new = new.replace(k, v)
        if full != new and para.runs:
            para.runs[0].text = new
            for r in para.runs[1:]:
                r.text = ""

    def _supprimer_nb(self, doc):
        a_supprimer = []
        for para in doc.paragraphs:
            if "NB pour les" in para.text:
                a_supprimer.append(para)
                continue
            # Les notes d'instruction du modele sont surlignees (couleur jaune) :
            # on les supprime aussi, meme si elles ne contiennent pas "NB pour les"
            if any(r.font.highlight_color not in (None, WD_COLOR_INDEX.AUTO)
                   for r in para.runs):
                a_supprimer.append(para)
        for para in a_supprimer:
            try:
                para._element.getparent().remove(para._element)
            except Exception:
                pass

    def _generer_word(self, modele, infos, png_path):
        plugin_dir   = os.path.dirname(__file__)
        template     = os.path.join(plugin_dir, "templates", modele + ".docx")
        doc          = Document(template)
        date_today   = datetime.now().strftime("%d/%m/%Y")
        commune      = infos["commune"]  or "Casablanca"
        province     = infos["province"] or "Casablanca"
        du_nom       = infos["du_nom"]   or "Plan d'Amenagement"
        du_type      = infos["du_type"]  or "Plan d'Amenagement"
        zones_txt    = ", ".join(
            z["zone"] + " secteur " + z["secteur"] for z in infos["zones"]) or "---"
        secteurs_txt = ", ".join(
            set(z["secteur"] for z in infos["zones"] if z["secteur"])) or "---"
        equip_txt    = ", ".join(infos["equipements"]) or "Aucun"

        subs = {
            "[Date d'elaboration de la NRU]":            date_today,
            "[Date d\u2019\u00e9laboration de la NRU]":  date_today,
            "[Date d'\u00e9laboration de la NRU]":       date_today,
            "[Date d'élaboration de la NRU]":            date_today,
            "[Nom du demandeur]":                         self.e_nom.text(),
            "[Adresse du demandeur]":                     self.e_adresse.text(),
            "[Reference fonciere du Terrain]":            self.e_ref.text(),
            "[R\u00e9f\u00e9rence fonci\u00e8re du Terrain]": self.e_ref.text(),
            "[Référence foncière du Terrain]":            self.e_ref.text(),
            "[Nom de la Commune]":                        commune,
            "[Nom de la Prefecture/Province]":            province,
            "[Nom de la Pr\u00e9fecture/Province]":       province,
            "[Nom de la Préfecture/Province]":            province,
            "[N de la NRU]":                              self.e_num.text(),
            "[N\u00b0 de la NRU]":                        self.e_num.text(),
            "[N° de la NRU]":                             self.e_num.text(),
            "[Date de la demande de la NRU]":             self.e_date.text(),
            "[Nom du DU]":                                du_nom,
            "[Type Document d'Urbanisme couvrant la parcelle (Plan d'am\u00e9nagement ou Plan de zonage ou Plan de d\u00e9veloppement d'agglom\u00e9ration rurale)]": du_type,
            "[Type Document d\u2019Urbanisme couvrant la parcelle (Plan d\u2019am\u00e9nagement ou Plan de zonage ou Plan de d\u00e9veloppement d\u2019agglom\u00e9ration rurale)]": du_type,
            "[Type Document d'Urbanisme (Plan d'am\u00e9nagement ou Plan de zonage ou Plan de d\u00e9veloppement d'agglom\u00e9ration rurale)]": du_type,
            "[N\u00b0 DU]":                               "N/A",
            "[N° DU]":                                    "N/A",
            "[Date d'approbation du DU]":                 "N/A",
            "[Date d\u2019approbation du DU]":            "N/A",
            "[N\u00b0 BO]":                               "N/A",
            "[N° BO]":                                    "N/A",
            "[Date de la BO]":                            "N/A",
            # NB : dans le modele Word, "Situe en zone :", "En secteur :" et
            # "Greve par :" sont des puces automatiques -> le tiret "- " ne fait
            # PAS partie du texte du paragraphe. On garde quand meme les variantes
            # avec tiret au cas ou un des 3 modeles (NRU-1/2/3) serait redige en texte simple.
            "- Situ\u00e9 en zone :":                     "- Situe en zone : " + zones_txt,
            "- Situé en zone :":                          "- Situe en zone : " + zones_txt,
            "Situ\u00e9 en zone :":                       "Situe en zone : " + zones_txt,
            "Situé en zone :":                            "Situe en zone : " + zones_txt,
            "- En secteur :":                             "- En secteur : " + secteurs_txt,
            "En secteur :":                                "En secteur : " + secteurs_txt,
            "- Grev\u00e9 par :":                         "- Greve par : " + equip_txt,
            "- Grevé par :":                              "- Greve par : " + equip_txt,
            "Grev\u00e9 par :":                           "Greve par : " + equip_txt,
            "Grevé par :":                                "Greve par : " + equip_txt,
        }

        for para in doc.paragraphs:
            self._remplacer_para(para, subs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        self._remplacer_para(para, subs)

        self._supprimer_nb(doc)

        doc.add_paragraph("")
        titre = doc.add_paragraph("PLAN DE SITUATION")
        if titre.runs:
            titre.runs[0].bold = True
        titre.alignment = 1
        if os.path.exists(png_path) and os.path.getsize(png_path) > 0:
            p_img = doc.add_paragraph()
            p_img.alignment = 1
            p_img.add_run().add_picture(png_path, width=Inches(5.5))
        else:
            doc.add_paragraph("[Plan de situation non disponible]")

        nom = ("NRU_" + self.e_nom.text().replace(" ", "_") +
               "_" + datetime.now().strftime("%Y%m%d_%H%M") + ".docx")
        out = os.path.join(self.e_sortie.text(), nom)
        doc.save(out)
        return out

    def generer(self):
        try:
            if not self.e_nom.text().strip():
                QMessageBox.warning(self, "Erreur", "Le nom du demandeur est obligatoire.")
                return
            if not self.e_gpkg.text().strip():
                QMessageBox.warning(self, "Erreur", "Selectionnez le fichier Casa_urba.gpkg.")
                return
            if not self.e_sortie.text().strip():
                QMessageBox.warning(self, "Erreur", "Selectionnez un dossier de sortie.")
                return

            self.progress.setValue(10)
            self._log("Chargement des couches GeoPackage...")
            gpkg = self.e_gpkg.text()
            couches = {
                "limite_du":   QgsVectorLayer(gpkg + "|layername=Limite_DU",  "LimiteDU",  "ogr"),
                "zones":       QgsVectorLayer(gpkg + "|layername=zones",       "Zones",     "ogr"),
                "communes":    QgsVectorLayer(gpkg + "|layername=communes",    "Communes",  "ogr"),
                "equipements": QgsVectorLayer(gpkg + "|layername=equipements", "Equip",     "ogr"),
            }

            self.progress.setValue(25)
            self._log("Construction de la geometrie de la parcelle...")
            geom = self._get_geom()

            self.progress.setValue(45)
            self._log("Analyse du logigramme NRU...")
            modele = self._get_modele(geom, couches["limite_du"])
            self._log("  --> Modele selectionne : " + modele)

            self.progress.setValue(60)
            self._log("Extraction des informations urbanistiques...")
            infos = self._extraire_infos(geom, couches)
            self._log("  --> Commune : " + (infos["commune"] or "non trouvee"))
            self._log("  --> Zones   : " + str(len(infos["zones"])) + " zone(s)")

            self.progress.setValue(75)
            self._log("Generation du plan de situation...")
            tmp_dir  = tempfile.mkdtemp(prefix="nru_")
            png_path = os.path.join(tmp_dir, "plan.png")
            self._export_png(geom, couches, png_path)

            self.progress.setValue(90)
            self._log("Generation du document Word...")
            word_path = self._generer_word(modele, infos, png_path)

            self.progress.setValue(100)
            self._log("Termine ! Fichier : " + word_path)
            QMessageBox.information(
                self, "NRU generee avec succes",
                "Modele utilise : " + modele +
                "\nCommune : " + (infos["commune"] or "N/A") +
                "\n\nFichier enregistre :\n" + word_path
            )

        except Exception as e:
            self.progress.setValue(0)
            QMessageBox.critical(self, "Erreur", str(e))
            self._log("ERREUR : " + str(e))